"""Anwendertexte: kein EXE-Kopieren, Modus 3 zweistufig, Foundry zentral."""

from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]


def test_anleitung_kein_exe_copy_und_modus3_pflicht():
    text = (ROOT / "ANLEITUNG.txt").read_text(encoding="utf-8").lower()
    assert "nicht auf den pc kopiert" in text or "nicht" in text and "kopiert" in text
    assert ".exe" not in text
    assert "ohne auswahl startet nichts" in text
    assert "foundry" in text
    assert "pseudokrat bleibt lokal" in text
    assert "pruefen.bat" in text
    assert "fachliche unterlagen" in text
    assert "ohne bestaetigung" in text


def test_installieren_kopiert_keine_exe():
    text = (ROOT / "Installieren.bat").read_text(encoding="utf-8", errors="replace").lower()
    assert "nicht auf den pc kopiert" in text
    assert "anhangspruefer.exe" not in text
    assert "starten.bat" in text


def test_struktur_start_ist_server_nicht_exe():
    text = (ROOT / "STRUKTUR.md").read_text(encoding="utf-8").lower()
    assert "start für anwender" in text
    assert "anwender kopieren keine exe" in text
    assert "llp_ai" in text


def test_vorstellung_kein_pc_copy_keine_exe():
    import docx

    doc = docx.Document(str(ROOT / "LLP Anhangspruefer - Vorstellung und Anleitung.docx"))
    blob = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            blob += "\n" + "\n".join(c.text for c in row.cells)
    low = blob.lower()
    assert "nicht auf den pc kopiert" in low
    assert "auf ihren pc kopiert" not in low
    assert "eine exe. doppelklick" not in low
    assert "keine exe auf den pc kopieren" in low
    assert "foundry" in low
    assert "n. a. (rechtsgrund)" in low
    assert "lokales sprachmodell" not in low


def test_werkzeuge_pseudokrat_ohne_foundry():
    text = (ROOT / "_Gemeinsam" / "WERKZEUGE.txt").read_text(encoding="utf-8").lower()
    assert "pseudokrat" in text
    assert "keine" in text and "foundry" in text
    assert "llp_ai" in text
