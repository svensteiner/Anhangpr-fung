# -*- coding: utf-8 -*-
"""Modus 2: Anhang als Word, nicht nur PDF."""

from __future__ import annotations

import io
import sys
from pathlib import Path

import docx

from anhangspruefer.pruefung.extractor import extract_from_anhang


def _anhang_docx(path: Path) -> None:
    doc = docx.Document()
    doc.add_paragraph("Haftungsverhältnisse")
    doc.add_paragraph("Gesamtbetrag der Haftungsverhältnisse 10.000,00 8.000,00")
    doc.add_paragraph("Durchschnittliche Zahl der Arbeitnehmer")
    doc.add_paragraph("Arbeiter 6 7")
    doc.add_paragraph("Angestellte 200 216")
    doc.add_paragraph("Gesamt 206 223")
    doc.save(str(path))


def test_extract_from_anhang_reads_docx(tmp_path):
    p = tmp_path / "anhang.docx"
    _anhang_docx(p)
    positions = extract_from_anhang(p)
    sections = {pos.section for pos in positions}
    assert "Haftungsverhaeltnisse" in sections
    haft = next(pos for pos in positions if pos.section == "Haftungsverhaeltnisse")
    assert haft.current_value == 10000.0
    assert haft.prior_value == 8000.0


def test_compare_route_rejects_wrong_suffix():
    root = Path(__file__).resolve().parents[2]
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    import app as webapp

    client = webapp.app.test_client()
    missing = client.post("/compare")
    assert missing.status_code == 400
    assert "PDF oder Word" in missing.get_json()["error"]

    wrong = client.post(
        "/compare",
        data={
            "current": (io.BytesIO(b"a"), "jetzt.txt"),
            "prior": (io.BytesIO(b"b"), "vorher.txt"),
        },
        content_type="multipart/form-data",
    )
    assert wrong.status_code == 400
    assert "Word" in wrong.get_json()["error"]


def test_pruefen_route_accepts_docx_and_rejects_other():
    root = Path(__file__).resolve().parents[2]
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    import app as webapp

    client = webapp.app.test_client()
    missing = client.post("/pruefen")
    assert missing.status_code == 400
    assert "Anhang" in missing.get_json()["error"]

    wrong = client.post(
        "/pruefen",
        data={"anhang": (io.BytesIO(b"nope"), "anhang.txt")},
        content_type="multipart/form-data",
    )
    assert wrong.status_code == 400
    assert "Word" in wrong.get_json()["error"]

    import tempfile

    with tempfile.TemporaryDirectory() as tmpdir:
        p = Path(tmpdir) / "anhang.docx"
        _anhang_docx(p)
        data = client.post(
            "/pruefen",
            data={"anhang": (p.open("rb"), "anhang.docx")},
            content_type="multipart/form-data",
        )
    assert data.status_code == 200, data.get_json()
    body = data.get_json()
    assert "filename" in body
    assert body["gesamt"] >= 1
